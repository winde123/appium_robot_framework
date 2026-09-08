#!/usr/bin/env python3
"""Build illustrated Word documents and an HTML gallery from captured evidence.

Requires python-docx and Pillow. Run after editing each category's coverage.json.
"""
import hashlib
import html
import json
import sys
import zipfile
from pathlib import Path
from xml.etree import ElementTree as ET

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor
from PIL import Image

ROOT = Path(__file__).resolve().parents[1] / "android-sgac2-2026-09-07"


def field(paragraph, instruction):
    element = OxmlElement("w:fldSimple")
    element.set(qn("w:instr"), instruction)
    paragraph._p.append(element)


def build(category):
    records = json.loads((category / "manifest.json").read_text())
    info = json.loads((category / "coverage.json").read_text())
    doc = Document()
    section = doc.sections[0]
    section.page_width, section.page_height = Inches(8.27), Inches(11.69)
    section.top_margin = section.bottom_margin = Inches(0.6)
    section.left_margin = section.right_margin = Inches(0.65)
    section.header_distance = section.footer_distance = Inches(0.25)
    normal = doc.styles["Normal"]
    normal.font.name = "Calibri"
    normal.font.size = Pt(10)
    normal.paragraph_format.space_after = Pt(5)
    for name in ("Title", "Heading 1", "Heading 2"):
        doc.styles[name].font.color.rgb = RGBColor.from_string("17224D")
    section.header.paragraphs[0].text = "MYICA MOBILE  |  ANDROID WALKTHROUGH  |  07 SEP 2026"
    section.header.paragraphs[0].style = "Caption"
    footer = section.footer.paragraphs[0]
    footer.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    footer.add_run("Staging build 2.0.0 / 420  •  Page ")
    field(footer, "PAGE")
    footer.add_run(" of ")
    field(footer, "NUMPAGES")
    doc.add_paragraph("PROJECT DOCUMENTATION", "Subtitle")
    doc.add_heading(info["title"], 0)
    doc.add_paragraph(info["summary"])
    doc.add_paragraph(f"{len(records)} screen captures • Pixel 7 Pro emulator • Android 16 • English")
    doc.add_paragraph("Captured from the installed staging app on 7 September 2026. Screenshots are original device PNGs. This is a screen walkthrough; it does not certify every test case or every input combination.")
    doc.add_heading("Coverage and outcomes", 1)
    table = doc.add_table(rows=1, cols=2)
    table.style = "Light Shading Accent 1"
    table.rows[0].cells[0].text = "Flow / branch"
    table.rows[0].cells[1].text = "Observed outcome"
    for flow in info["flows"]:
        cells = table.add_row().cells
        cells[0].text = flow["name"]
        cells[1].text = flow["outcome"]
    for note in info.get("notes", []):
        doc.add_paragraph(note)
    doc.add_heading("How to use this document", 1)
    doc.add_paragraph("Each following page records one screen, the action used to reach it, and any relevant observation. Screen numbers match the category manifest and PNG filenames. Source XML and SHA-256 hashes are retained beside the screenshots for traceability. Open the original PNG when you need to zoom into small text.")
    for r in records:
        path = category / r["screenshot"]
        assert hashlib.sha256(path.read_bytes()).hexdigest() == r["sha256"], path
        with Image.open(path) as im:
            width_px, height_px = im.size
            im.verify()
        ET.parse(category / r["source"])
        heading = doc.add_heading(f"{r['step']:03d}  {r['title']}", 1)
        heading.paragraph_format.page_break_before = True
        if r["action"]:
            doc.add_paragraph(r["action"])
        if r["observation"]:
            doc.add_paragraph(r["observation"])
        if r["status"] != "Captured":
            doc.add_paragraph("Outcome: " + r["status"])
        p = doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.paragraph_format.space_after = Pt(0)
        p.paragraph_format.keep_with_next = True
        display_height = min(6.8, 6.7 * height_px / width_px)
        p.add_run().add_picture(str(path), height=Inches(display_height))
        cap = doc.add_paragraph(f"Screen {r['step']:03d} • {r['captured_at'][:19].replace('T', ' ')} • {r['status']}", "Caption")
        cap.alignment = WD_ALIGN_PARAGRAPH.CENTER
    target = category / (category.name + ".docx")
    doc.save(target)
    with zipfile.ZipFile(target) as archive:
        assert archive.testzip() is None
        drawing_count = len(ET.fromstring(archive.read("word/document.xml")).findall(".//{http://schemas.openxmlformats.org/wordprocessingml/2006/main}drawing"))
        assert drawing_count == len(records), (drawing_count, len(records))
    md = [f"# {info['title']}", "", "Last reviewed: 2026-09-07", "", info["summary"], "", f"[Word walkthrough]({target.name})", "", "| Flow | Observed outcome |", "| --- | --- |"]
    md += [f"| {f['name']} | {f['outcome']} |" for f in info["flows"]]
    md += ["", "## Screens", "", "| Step | Screen | Action | Outcome |", "| --- | --- | --- | --- |"]
    md += [f"| {r['step']:03d} | [{r['title']}]({r['screenshot']}) | {r['action']} | {r['observation'] or r['status']} |" for r in records]
    (category / "README.md").write_text("\n".join(md)+"\n")
    return records, info, target


def main():
    summaries = []
    gallery = ["<!doctype html><html lang='en'><meta charset='utf-8'><meta name='viewport' content='width=device-width,initial-scale=1'><title>MyICA project walkthrough</title><style>body{font:16px system-ui;background:#f4f6fa;color:#17224d;margin:2rem}a{color:#2156a5}nav{display:flex;gap:1rem;flex-wrap:wrap}.grid{display:grid;grid-template-columns:repeat(auto-fill,minmax(260px,1fr));gap:24px}article{padding:20px;background:white;border-radius:12px}img{width:100%;height:auto}h2{margin-top:3rem}p{line-height:1.5}</style><h1>MyICA Mobile: Android screen walkthrough</h1><p>7 September 2026 • Staging 2.0.0 / 420 • Pixel 7 Pro</p><nav>"]
    results = []
    for manifest in sorted(ROOT.glob("*/manifest.json")):
        results.append((manifest.parent, *build(manifest.parent)))
    for category, records, info, target in results:
        gallery.append(f"<a href='#{category.name}'>{html.escape(info['title'])}</a>")
    gallery.append("</nav>")
    for category, records, info, target in results:
        gallery.append(f"<h2 id='{category.name}'>{html.escape(info['title'])}</h2><p>{html.escape(info['summary'])}</p><p><a href='{category.name}/{target.name}'>Download Word document</a></p><ul>")
        for f in info["flows"]:
            gallery.append(f"<li><strong>{html.escape(f['name'])}:</strong> {html.escape(f['outcome'])}</li>")
        gallery.append("</ul><div class='grid'>")
        for r in records:
            p = f"{category.name}/{r['screenshot']}"
            gallery.append(f"<article><h3>{r['step']:03d} {html.escape(r['title'])}</h3><p>{html.escape(r['action'])}</p><a href='{p}'><img loading='lazy' src='{p}' alt='{html.escape(r['title'], quote=True)}'></a><p>{html.escape(r['observation'] or r['status'])}</p></article>")
        gallery.append("</div>")
        summaries.append({"category":category.name,"screens":len(records),"docx":str(target.relative_to(ROOT)),"bytes":target.stat().st_size,"verified":True})
    gallery.append("</html>")
    (ROOT / "index.html").write_text("\n".join(gallery))
    (ROOT / "validation.json").write_text(json.dumps(summaries, indent=2)+"\n")
    print(json.dumps(summaries, indent=2))


if __name__ == "__main__":
    main()
